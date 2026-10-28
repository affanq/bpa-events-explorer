"""Read uploaded submissions (PDF, DOCX, PPTX, TXT/MD) into text; nothing is stored."""

import io
import re
import zipfile
from pathlib import Path
from typing import Protocol

import pymupdf
import streamlit as st
from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError

from bpa_app import config

SUPPORTED_TYPES = ("pdf", "docx", "pptx", "txt", "md")
TRUNCATION_NOTE = "[Text truncated: the rest of the document was not read.]"

_BLANK_RUN = re.compile(r"\n{3,}")
_OFFICE_ERRORS = (
    zipfile.BadZipFile,
    DocxPackageNotFoundError,
    PptxPackageNotFoundError,
    KeyError,
    ValueError,
)


class UploadError(Exception):
    """A problem with the uploaded file; the message is safe to show to users."""


class UploadedFileLike(Protocol):
    """The parts of Streamlit's UploadedFile that are used here."""

    name: str

    def getvalue(self) -> bytes: ...


def _check_size(size: int | None) -> None:
    if size is not None and size > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise UploadError(
            f"This file is {size / 1024 / 1024:.1f} MB; "
            f"the limit is {config.MAX_UPLOAD_MB} MB."
        )


def _table_lines(rows: list[list[str]]) -> list[str]:
    lines = [" | ".join(" ".join(cell.split()) for cell in row) for row in rows]
    return [line for line in lines if line.strip(" |")]


def _read_pdf(data: bytes) -> str:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise UploadError("This PDF is password-protected.")
            return "\n\n".join(page.get_text() for page in doc)
    except (pymupdf.FileDataError, RuntimeError) as exc:
        raise UploadError("Could not read this PDF. The file may be damaged.") from exc


def _read_docx(data: bytes) -> str:
    try:
        parts: list[str] = []
        for block in Document(io.BytesIO(data)).iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                rows = []
                for row in block.rows:
                    seen: set[int] = set()  # merged cells repeat the same cell
                    cells = []
                    for cell in row.cells:
                        if id(cell._tc) not in seen:
                            seen.add(id(cell._tc))
                            cells.append(cell.text)
                    rows.append(cells)
                parts.extend(_table_lines(rows))
    except _OFFICE_ERRORS as exc:
        raise UploadError("Could not read this Word file. Is it a .docx?") from exc
    return "\n".join(parts)


def _shape_text(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [text for child in shape.shapes for text in _shape_text(child)]
    if shape.has_table:
        rows = [
            [cell.text for cell in row.cells if not cell.is_spanned]
            for row in shape.table.rows
        ]
        return _table_lines(rows)
    if shape.has_text_frame and shape.text_frame.text.strip():
        return [shape.text_frame.text]
    return []


def _read_pptx(data: bytes) -> str:
    try:
        slides = []
        for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
            lines = [text for shape in slide.shapes for text in _shape_text(shape)]
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"Notes: {notes}")
            if lines:
                slides.append(f"Slide {number}\n" + "\n".join(lines))
    except _OFFICE_ERRORS as exc:
        raise UploadError(
            "Could not read this PowerPoint file. Is it a .pptx?"
        ) from exc
    return "\n\n".join(slides)


def _read_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


_READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    ".txt": _read_text,
    ".md": _read_text,
}


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANK_RUN.sub("\n\n", text).strip()


def _truncate(text: str) -> str:
    limit = config.MAX_UPLOAD_TEXT_CHARS
    if len(text) <= limit:
        return text
    st.warning(
        f"This file is long, so only the first {limit:,} characters will be used."
    )
    return f"{text[:limit].rstrip()}\n\n{TRUNCATION_NOTE}"


def read_upload(uploaded_file: UploadedFileLike) -> tuple[str, bytes | None]:
    """Return (text, pdf_bytes); pdf_bytes is None unless the file is a PDF.

    Raises UploadError for oversized, unsupported or unreadable files. Long text is cut
    to MAX_UPLOAD_TEXT_CHARS with a Streamlit warning and a note at the end of the text.
    """
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in _READERS:
        raise UploadError(
            "Unsupported file type. Please upload a "
            + ", ".join(f".{t}" for t in SUPPORTED_TYPES)
            + " file."
        )
    _check_size(getattr(uploaded_file, "size", None))
    data = uploaded_file.getvalue()
    _check_size(len(data))

    text = _truncate(_clean(_READERS[suffix](data)))
    if suffix == ".pdf":
        if not text:
            st.warning("No text found in this PDF; it may be a scanned image.")
        return text, data
    if not text:
        raise UploadError("No readable text was found in this file.")
    return text, None
