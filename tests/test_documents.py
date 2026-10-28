import io
from dataclasses import dataclass

import pymupdf
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from bpa_app import config, documents
from bpa_app.documents import UploadError, read_upload


@dataclass
class FakeUpload:
    name: str
    data: bytes
    size: int | None = None

    def getvalue(self) -> bytes:
        return self.data


@pytest.fixture
def warnings_seen(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    seen: list[str] = []
    monkeypatch.setattr(documents.st, "warning", seen.append)
    return seen


def make_pdf(*pages: str) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    return doc.tobytes()


def make_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Intro paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Criterion"
    table.cell(0, 1).text = "Points"
    table.cell(1, 0).text = "Creativity"
    table.cell(1, 1).text = "20"
    doc.add_paragraph("Closing paragraph")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_pptx() -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    first = prs.slides.add_slide(blank)
    first.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.text = "Welcome"
    group = first.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(3), Inches(4), Inches(1)
    ).text_frame.text = "Grouped"
    first.notes_slide.notes_text_frame.text = "Say hello"

    prs.slides.add_slide(blank)  # no text: should be skipped

    third = prs.slides.add_slide(blank)
    table = third.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(1)
    ).table
    for row, values in enumerate((("Item", "Cost"), ("Laptop", "$900"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_pdf_returns_text_and_original_bytes() -> None:
    data = make_pdf("First page", "Second page")
    text, pdf = read_upload(FakeUpload("Entry.PDF", data))
    assert text.index("First page") < text.index("Second page")
    assert pdf == data


def test_pdf_without_text_is_allowed_but_warns(warnings_seen: list[str]) -> None:
    data = make_pdf("")
    assert read_upload(FakeUpload("scan.pdf", data)) == ("", data)
    assert "scanned" in warnings_seen[0]


def test_docx_reads_paragraphs_and_tables_in_order() -> None:
    text, pdf = read_upload(FakeUpload("plan.docx", make_docx()))
    assert pdf is None
    assert (
        text
        == "Intro paragraph\nCriterion | Points\nCreativity | 20\nClosing paragraph"
    )


def test_pptx_reads_all_shapes_notes_and_tables_with_slide_prefix() -> None:
    text, pdf = read_upload(FakeUpload("deck.pptx", make_pptx()))
    assert pdf is None
    assert text == (
        "Slide 1\nWelcome\nGrouped\nNotes: Say hello\n\n"
        "Slide 3\nItem | Cost\nLaptop | $900"
    )


@pytest.mark.parametrize("name", ["notes.txt", "notes.md"])
def test_text_files_decode_utf8_with_or_without_bom(name: str) -> None:
    assert read_upload(FakeUpload(name, "# Title\r\n\r\n\r\n\r\nCafé".encode()))[0] == (
        "# Title\n\nCafé"
    )
    assert read_upload(FakeUpload(name, b"\xef\xbb\xbfHello"))[0] == "Hello"


def test_text_falls_back_to_windows_encoding() -> None:
    assert read_upload(FakeUpload("old.txt", "Café".encode("cp1252")))[0] == "Café"


def test_long_text_is_truncated_with_warning(
    monkeypatch: pytest.MonkeyPatch, warnings_seen: list[str]
) -> None:
    monkeypatch.setattr(config, "MAX_UPLOAD_TEXT_CHARS", 100)
    text, _ = read_upload(FakeUpload("big.txt", ("word " * 100).encode()))
    assert text.startswith("word word") and text.endswith(documents.TRUNCATION_NOTE)
    assert len(text) < 100 + len(documents.TRUNCATION_NOTE) + 3
    assert "100" in warnings_seen[0]


def test_default_truncation_limit_is_about_60k_characters() -> None:
    assert config.MAX_UPLOAD_TEXT_CHARS == 60_000


def test_text_at_the_limit_is_not_truncated(
    monkeypatch: pytest.MonkeyPatch, warnings_seen: list[str]
) -> None:
    monkeypatch.setattr(config, "MAX_UPLOAD_TEXT_CHARS", 10)
    assert read_upload(FakeUpload("ok.txt", b"0123456789"))[0] == "0123456789"
    assert warnings_seen == []


def test_oversized_upload_is_rejected_before_reading() -> None:
    too_big = (config.MAX_UPLOAD_MB + 1) * 1024 * 1024

    class Unreadable(FakeUpload):
        def getvalue(self) -> bytes:
            raise AssertionError("must not read an oversized file")

    with pytest.raises(UploadError, match=r"limit is 10 MB"):
        read_upload(Unreadable("big.pdf", b"", size=too_big))


def test_size_is_also_checked_on_the_bytes_when_size_is_unknown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(config, "MAX_UPLOAD_MB", 0.0001)
    with pytest.raises(UploadError, match="limit is"):
        read_upload(FakeUpload("a.txt", b"x" * 1000))


@pytest.mark.parametrize("name", ["virus.exe", "old.doc", "slides.ppt", "noextension"])
def test_unsupported_types_are_rejected(name: str) -> None:
    with pytest.raises(UploadError, match="Unsupported file type"):
        read_upload(FakeUpload(name, b"data"))


@pytest.mark.parametrize(
    ("name", "message"),
    [
        ("bad.pdf", "PDF"),
        ("bad.docx", "Word"),
        ("bad.pptx", "PowerPoint"),
    ],
)
def test_corrupt_files_give_friendly_errors(name: str, message: str) -> None:
    with pytest.raises(UploadError, match=message):
        read_upload(FakeUpload(name, b"this is not a real file"))


def test_a_zip_that_is_not_an_office_file_is_rejected() -> None:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr("hello.txt", "hi")
    for name in ("fake.docx", "fake.pptx"):
        with pytest.raises(UploadError):
            read_upload(FakeUpload(name, out.getvalue()))


def test_password_protected_pdf_is_rejected() -> None:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    data = doc.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw"
    )
    with pytest.raises(UploadError, match="password"):
        read_upload(FakeUpload("locked.pdf", data))


def test_empty_non_pdf_is_rejected() -> None:
    with pytest.raises(UploadError, match="No readable text"):
        read_upload(FakeUpload("empty.txt", b"  \n\n "))
    empty_docx = io.BytesIO()
    Document().save(empty_docx)
    with pytest.raises(UploadError, match="No readable text"):
        read_upload(FakeUpload("empty.docx", empty_docx.getvalue()))
